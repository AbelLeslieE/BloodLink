from datetime import datetime, timedelta, timezone
from io import BytesIO
import base64
import secrets
import sqlite3
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
import jwt
from jwt import InvalidTokenError as JWTError
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from backend.auth.security import (create_access_token, create_password_reset_token,
    get_token_subject, get_password_reset_data, hash_password, verify_password)
from backend.config.settings import get_settings, ConfigurationError
from backend.database.models import Donor, User, BloodRequest
from backend.security.database import configure_sqlcipher, create_database_engine
from backend.security.encrypt_database import encrypted_copy
from backend.security.exports import protect_workbook, safe_spreadsheet_cell
from backend.security.push import validate_push_endpoint
from backend.security.uploads import validate_workbook_archive


def test_token_purpose_expiration_and_tampering():
    assert get_token_subject(create_access_token("admin", 0)) == "admin"
    with pytest.raises(JWTError):
        get_token_subject(create_password_reset_token("admin", 0))
    with pytest.raises(JWTError):
        get_password_reset_data(create_access_token("admin", 0))
    settings = get_settings()
    for payload in ({"sub": "admin", "purpose": "access", "ver": 0},
                    {"sub": "admin", "purpose": "access", "ver": 0, "exp": datetime.now(timezone.utc)-timedelta(seconds=1)}):
        with pytest.raises(JWTError):
            get_token_subject(jwt.encode(payload, settings.secret_key, algorithm="HS256"))


def test_long_passwords_not_silently_truncated():
    first = "a" * 72 + "first"
    hashed = hash_password(first)
    assert verify_password(first, hashed)
    assert not verify_password("a" * 72 + "second", hashed)


@pytest.mark.parametrize("path", ["/api/donations/", "/api/donations/summary", "/api/donations/recent",
    "/api/donors", "/api/users", "/api/dashboard", "/api/blood-requests", "/api/notifications"])
def test_admin_routes_reject_anonymous_and_donors(system, path):
    client, _, tokens = system
    assert client.get(path).status_code == 401
    assert client.get(path, headers=tokens["donor1"]).status_code == 403
    assert client.get(path, headers=tokens["admin"]).status_code == 200


def test_reset_token_cannot_access_api(system):
    client, _, _ = system
    response = client.get("/api/auth/me", headers={"Authorization": "Bearer " + create_password_reset_token("admin", 0)})
    assert response.status_code == 401


def test_cookie_login_csrf_logout_and_revocation(system):
    client, _, tokens = system
    response = client.post("/api/auth/login", data={"username": "admin", "password": "AdminPassword123"}, headers={"X-Session-Mode": "cookie"})
    assert response.status_code == 200
    assert response.json()["access_token"] == "cookie-session"
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=strict" in response.headers["set-cookie"]
    assert client.get("/api/auth/me").status_code == 200
    assert client.post("/api/auth/logout", headers={"Origin": "https://attacker.example"}).status_code == 403
    assert client.post("/api/auth/logout", headers={"Origin": "http://testserver"}).status_code == 200
    assert client.get("/api/auth/me").status_code == 401
    assert client.get("/api/auth/me", headers=tokens["admin"]).status_code == 401


def test_security_headers_and_validation_redaction(system):
    client, _, _ = system
    response = client.get("/login")
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["referrer-policy"] == "no-referrer"
    invalid = client.post("/api/auth/password-reset/confirm", json={"token": "private-token", "new_password": "private-password"})
    assert invalid.status_code == 422
    assert "private-token" not in invalid.text
    assert "private-password" not in invalid.text
    assert client.post("/api/auth/login", content=b"x" * (256*1024+1)).status_code == 413


def test_shared_rate_limit(system):
    client, _, _ = system
    # Invalid input still consumes the public endpoint's budget, with no email.
    for _ in range(30):
        assert client.post("/api/auth/password-reset/request", json={}).status_code == 422
    response = client.post("/api/auth/password-reset/request", json={})
    assert response.status_code == 429 and response.headers["retry-after"] == "900"


@pytest.mark.parametrize("url", ["http://fcm.googleapis.com/a", "https://127.0.0.1/a", "https://169.254.169.254/a",
    "https://fcm.googleapis.com.attacker.test/a", "https://fcm.googleapis.com:8443/a", "https://user@fcm.googleapis.com/a"])
def test_push_ssrf_rejection(url):
    with pytest.raises(ValueError):
        validate_push_endpoint(url)


def test_push_provider_acceptance():
    for url in ("https://fcm.googleapis.com/fcm/send/abc", "https://updates.push.services.mozilla.com/wpush/v2/abc",
                "https://web.push.apple.com/abc", "https://wns1.notify.windows.com/abc"):
        assert validate_push_endpoint(url) == url


def test_push_subscription_cannot_be_stolen(system):
    client, _, tokens = system
    payload = {"endpoint": "https://fcm.googleapis.com/fcm/send/example", "keys": {
        "p256dh": base64.urlsafe_b64encode(b"\x04" + secrets.token_bytes(64)).decode().rstrip("="),
        "auth": base64.urlsafe_b64encode(secrets.token_bytes(16)).decode().rstrip("=")}}
    assert client.post("/api/push/subscriptions", json=payload, headers=tokens["donor1"]).status_code == 204
    assert client.post("/api/push/subscriptions", json=payload, headers=tokens["donor2"]).status_code == 409


def test_formula_injection_protection():
    workbook = Workbook()
    workbook.active.append(["=HYPERLINK(\"https://attacker.test\")", "+12345", "Normal name", 25])
    protect_workbook(workbook)
    output = BytesIO(); workbook.save(output); output.seek(0)
    sheet = load_workbook(output).active
    assert sheet["A1"].data_type == "s" and sheet["A1"].value.startswith("'")
    assert sheet["B1"].data_type == "s"
    assert sheet["D1"].value == 25
    assert safe_spreadsheet_cell("  =1+1").startswith("'")


def test_workbook_archive_limits():
    data = BytesIO()
    with ZipFile(data, "w") as archive:
        archive.writestr("[Content_Types].xml", "")
        archive.writestr("oversized.xml", b"a" * (21 * 1024 * 1024))
    with pytest.raises(ValueError):
        validate_workbook_archive(data.getvalue())


def test_sqlcipher_copy_and_wrong_key(tmp_path):
    from sqlcipher3 import dbapi2
    source, target = tmp_path / "plain.db", tmp_path / "encrypted.db"
    with sqlite3.connect(source) as db:
        db.execute("CREATE TABLE donors(id INTEGER PRIMARY KEY, name TEXT)")
        db.execute("INSERT INTO donors(name) VALUES (?)", ("Sensitive Test Donor",))
    original = source.read_bytes()
    key = secrets.token_hex(32)
    assert encrypted_copy(source, target, key) == {"donors": 1}
    assert source.read_bytes() == original
    assert b"Sensitive Test Donor" not in target.read_bytes()
    with pytest.raises(sqlite3.DatabaseError):
        with sqlite3.connect(target) as db:
            db.execute("SELECT * FROM donors").fetchall()
    db = dbapi2.connect(str(target))
    try:
        configure_sqlcipher(db, key)
        assert db.execute("SELECT name FROM donors WHERE name LIKE ?", ("%Test%",)).fetchone()[0] == "Sensitive Test Donor"
    finally:
        db.close()
    wrong = dbapi2.connect(str(target))
    try:
        with pytest.raises(dbapi2.DatabaseError):
            configure_sqlcipher(wrong, secrets.token_hex(32))
    finally:
        wrong.close()
    with pytest.raises(ValueError):
        encrypted_copy(source, target, key)


def test_production_plaintext_sqlite_refused():
    with pytest.raises(ConfigurationError):
        create_database_engine(replace(get_settings(), production=True))


def registration(username="newdonor", email="new@example.org", phone="9999900010"):
    return {"full_name": "New Donor", "email": email, "phone": phone, "username": username,
            "blood_group": "A+", "gender": "Female", "date_of_birth": "1995-01-20",
            "current_status": "Unemployed", "password": "StrongPassword123", "confirm_password": "StrongPassword123"}


def test_registration_preserved_but_disabled_account_takeover_blocked(system):
    client, sessions, tokens = system
    assert client.post("/api/donor-registration/complete", json=registration()).status_code == 201
    assert client.post("/api/auth/login", data={"username": "newdonor", "password": "StrongPassword123"}).status_code == 200
    visible_donors = client.get("/api/donors", headers=tokens["admin"])
    assert visible_donors.status_code == 200
    assert any(item["email"] == "new@example.org" for item in visible_donors.json())
    with sessions() as db:
        donor = db.scalar(select(Donor).where(Donor.email == "new@example.org"))
        user = db.scalar(select(User).where(User.username == "newdonor"))
        assert user.donor_id == donor.id
        assert donor.gender == "Female"
        assert donor.date_of_birth.isoformat() == "1995-01-20"
    with sessions.begin() as db:
        admin = db.scalar(select(User).where(User.username == "admin"))
        admin.active = False
        old_hash = admin.password_hash
    attack = client.post("/api/donor-registration/complete", json=registration("admin", "admin@example.org", "9999900000"))
    assert attack.status_code == 409
    with sessions() as db:
        admin = db.scalar(select(User).where(User.username == "admin"))
        assert not admin.active and admin.password_hash == old_hash


def test_user_editor_preserves_donor_link_and_rejects_orphans(system):
    client, sessions, tokens = system
    with sessions() as db:
        donor_user_id = db.scalar(select(User.id).where(User.username == "donor1"))
    update_payload = {
        "full_name": "Test Donor 1 Updated", "department": "NSS", "role": "Donor",
        "email": "donor1@example.org", "phone": "9999900001", "username": "donor1",
    }
    response = client.put(f"/api/users/{donor_user_id}", json=update_payload, headers=tokens["admin"])
    assert response.status_code == 200, response.text
    assert response.json()["user"]["donor_id"] == 1
    with sessions() as db:
        assert db.get(User, donor_user_id).donor_id == 1

    orphan_payload = {
        **update_payload, "full_name": "Orphan Donor", "email": "orphan@example.org",
        "phone": "9999900098", "username": "orphan", "password": "StrongPassword123",
    }
    rejected = client.post("/api/users/register", json=orphan_payload, headers=tokens["admin"])
    assert rejected.status_code == 422


def test_donor_portal_patient_privacy(system):
    client, sessions, tokens = system
    with sessions.begin() as db:
        db.add(BloodRequest(patient_name="Private Patient", case_details="Private Condition", blood_group="A+",
            units_required=1, required_date=datetime.now().date(), priority="Normal", hospital_name="Test Hospital",
            hospital_location="Test City", contact_person="Private Contact", contact_phone="9999900099", created_by=1))
    response = client.get("/api/donor-dashboard/requests", headers=tokens["donor1"])
    assert response.status_code == 200
    assert "Private Patient" not in response.text and "9999900099" not in response.text
    assert "Test Hospital" in response.text


def test_hashed_email_links_and_single_use(system):
    from backend.database import crud
    from backend.database.notification import Notification
    from backend.database.notification_recipient import NotificationRecipient
    from backend.database.donor_response import DonorResponse
    client, sessions, _ = system
    raw = secrets.token_urlsafe(32)
    with sessions() as db:
        request = BloodRequest(patient_name="Test", case_details="Test", blood_group="A+", units_required=1,
            required_date=datetime.now().date(), priority="Normal", hospital_name="Test", hospital_location="Test",
            contact_person="Test", contact_phone="9999900099", created_by=1)
        db.add(request); db.flush()
        notification = Notification(blood_request_id=request.id, title="Test", status="ACTIVE")
        db.add(notification); db.flush()
        recipient = NotificationRecipient(notification_id=notification.id, donor_id=1, email="donor1@example.org")
        db.add(recipient); db.flush()
        record = crud.create_email_token(db, recipient, raw, datetime.now(timezone.utc)+timedelta(hours=1))
        digest = record.token
        assert digest.startswith("sha256:") and raw not in digest
        assert crud.get_email_token(db, raw).id == record.id
        assert crud.get_email_token(db, digest) is None
    assert client.get("/email/accept/"+raw).status_code == 200
    assert client.post("/email/accept/"+raw).status_code == 200
    assert client.post("/email/decline/"+raw).status_code == 200
    with sessions() as db:
        assert db.scalar(select(DonorResponse)).response == "ACCEPTED"


def test_email_password_setup_remains_single_use(system):
    from backend.services.pending_registration_service import issue_password_setup_token
    client, sessions, _ = system
    data = registration()
    data.pop("password"); data.pop("confirm_password")
    assert client.post("/api/donor-registration/verify-details", json=data).status_code == 200
    with sessions() as db:
        user = db.scalar(select(User).where(User.username == "newdonor"))
        token = issue_password_setup_token(db, user)
    payload = {"token": token, "new_password": "SetupPassword123"}
    assert client.post("/api/donor-registration/password-setup/confirm", json=payload).status_code == 200
    assert client.post("/api/donor-registration/password-setup/confirm", json=payload).status_code == 400
    assert client.post("/api/auth/login", data={"username": "newdonor", "password": "SetupPassword123"}).status_code == 200


def test_access_log_redaction():
    import logging
    from backend.security.logging import RedactAccessLog
    record = logging.LogRecord("uvicorn.access", logging.INFO, "", 1, "%s - %s %s HTTP/%s %s",
        ("private-client-ip", "GET", "/email/accept/secret-link?token=secret-query", "1.1", 200), None)
    RedactAccessLog().filter(record)
    assert "secret" not in record.getMessage() and "private-client-ip" not in record.getMessage()


def test_push_redirects_disabled(monkeypatch):
    import requests
    from backend.security.push import PushSession
    captured = {}
    def fake_request(self, method, url, **kwargs):
        captured.update(kwargs)
    monkeypatch.setattr(requests.Session, "request", fake_request)
    with PushSession() as session:
        session.request("POST", "https://fcm.googleapis.com/fcm/send/test", allow_redirects=True)
    assert captured["allow_redirects"] is False and captured["timeout"] == 10


def test_cookie_secure_on_https_and_script_policy(system):
    client, _, _ = system
    response = client.post("https://testserver/api/auth/login", data={"username": "admin", "password": "AdminPassword123"}, headers={"X-Session-Mode": "cookie"})
    assert "; Secure" in response.headers["set-cookie"]
    assert "script-src 'self';" in response.headers["content-security-policy"]


def test_rate_limit_atomic_across_threads(system):
    from concurrent.futures import ThreadPoolExecutor
    from backend.security.rate_limit import consume_limit
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda _: consume_limit("concurrency-test", "test", 5, 60), range(20)))
    assert sum(results) == 5


def test_production_host_https_and_documentation_controls(tmp_path):
    import os
    import subprocess
    import sys
    env = dict(os.environ, APP_ENV="production", ALLOWED_HOSTS="testserver",
               BACKEND_URL="https://testserver", FRONTEND_URL="https://testserver",
               DATABASE_URL="sqlite:///"+str(tmp_path/"prod.db"), DATABASE_ENCRYPTION_KEY=secrets.token_hex(32))
    code = '''from fastapi.testclient import TestClient
from backend.main import app
c = TestClient(app)
assert c.get('http://testserver/login', follow_redirects=False).status_code == 307
r = c.get('https://testserver/login')
assert r.status_code == 200 and 'max-age=' in r.headers['strict-transport-security']
assert c.get('https://testserver/docs').status_code == 404
assert c.get('https://testserver/openapi.json').status_code == 404
assert c.get('https://attacker.test/login').status_code == 400
'''
    subprocess.run([sys.executable, "-c", code], env=env, check=True, capture_output=True, text=True)
