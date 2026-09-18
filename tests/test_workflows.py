from datetime import date, timedelta
from io import BytesIO
import os
import secrets
from alembic import command
from alembic.config import Config
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from backend.config.settings import get_settings
from backend.database.models import BloodRequest, Donor, User, DonationHistory
from backend.security.database import create_database_engine
from backend.auth.security import create_password_reset_token
from backend.services.donor_matching_service import (
    COMPATIBILITY_MATRIX,
    is_compatible_donor,
)


def blood_request():
    return {"patient_name": "Regression Patient", "case_details": "Regression only", "blood_group": "A+",
        "units_required": 1, "required_date": str(date.today()+timedelta(days=1)), "priority": "Normal",
        "hospital_name": "Regression Hospital", "hospital_location": "Test City",
        "contact_person": "Test Contact", "contact_phone": "9999900020"}


def test_red_cell_compatibility_matrix_is_directional():
    assert COMPATIBILITY_MATRIX == {
        "A+": ["A+", "A-", "O+", "O-"],
        "A-": ["A-", "O-"],
        "B+": ["B+", "B-", "O+", "O-"],
        "B-": ["B-", "O-"],
        "AB+": ["A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"],
        "AB-": ["A-", "B-", "AB-", "O-"],
        "O+": ["O+", "O-"],
        "O-": ["O-"],
    }
    assert is_compatible_donor("A+", "O-")
    assert is_compatible_donor(" ab+ ", " b- ")
    assert not is_compatible_donor("O-", "O+")
    assert not is_compatible_donor("A-", "B-")


def test_find_save_and_send_use_real_blood_compatibility(system):
    client, sessions, tokens = system
    with sessions.begin() as db:
        compatible = [
            Donor(donor_code="COMP-A-", full_name="A Negative", blood_group="A-",
                  phone="9999900051", email="a-negative@example.org", status="Available"),
            Donor(donor_code="COMP-O-", full_name="O Negative", blood_group="O-",
                  phone="9999900052", email="o-negative@example.org", status="Available"),
        ]
        incompatible = Donor(
            donor_code="COMP-B+", full_name="B Positive", blood_group="B+",
            phone="9999900053", email="b-positive@example.org", status="Available",
        )
        db.add_all([*compatible, incompatible])
        db.flush()
        compatible_ids = [donor.id for donor in compatible]
        incompatible_id = incompatible.id
        request = BloodRequest(
            patient_name="Compatibility Patient", case_details="Compatibility regression",
            blood_group="A-", units_required=1, required_date=date.today() + timedelta(days=1),
            priority="Urgent", hospital_name="Compatibility Hospital",
            hospital_location="Test City", contact_person="Test Contact",
            contact_phone="9999900054", created_by=1,
        )
        db.add(request)
        db.flush()
        request_id = request.id

    matched = client.post(
        "/api/match/find", json={"blood_request_id": request_id}, headers=tokens["admin"]
    )
    assert matched.status_code == 200, matched.text
    matches = matched.json()["matches"]
    assert {item["donor"]["blood_group"] for item in matches} == {"A-", "O-"}
    assert {item["compatibility_percent"] for item in matches} == {80, 100}
    assert all(0 <= item["compatibility_percent"] <= 100 for item in matches)

    saved = client.post(
        "/api/match/save",
        json={"blood_request_id": request_id, "donor_ids": compatible_ids},
        headers=tokens["admin"],
    )
    assert saved.status_code == 200 and saved.json()["saved_count"] == 2
    rejected = client.post(
        "/api/match/save",
        json={"blood_request_id": request_id, "donor_ids": [incompatible_id]},
        headers=tokens["admin"],
    )
    assert rejected.status_code == 409

    rejected_send = client.post(
        "/api/match/send",
        json={"blood_request_id": request_id, "donor_ids": [incompatible_id]},
        headers=tokens["admin"],
    )
    assert rejected_send.status_code == 409

    sent = client.post(
        "/api/match/send",
        json={"blood_request_id": request_id, "donor_ids": compatible_ids},
        headers=tokens["admin"],
    )
    assert sent.status_code == 200, sent.text


def test_compatible_request_reaches_donor_portal_and_confirmation(system):
    client, sessions, tokens = system
    with sessions.begin() as db:
        donor_id = db.scalar(select(Donor.id).where(Donor.email == "donor1@example.org"))
        request = BloodRequest(
            patient_name="AB Positive Patient", case_details="Compatibility regression",
            blood_group="AB+", units_required=1, required_date=date.today() + timedelta(days=1),
            priority="Urgent", hospital_name="Compatibility Hospital",
            hospital_location="Test City", contact_person="Test Contact",
            contact_phone="9999900055", created_by=1,
        )
        db.add(request)
        db.flush()
        request_id = request.id

    visible = client.get("/api/donor-dashboard/requests", headers=tokens["donor1"])
    assert visible.status_code == 200
    assert request_id in {item["id"] for item in visible.json()}
    response = client.post(
        f"/api/donor-dashboard/requests/{request_id}/response",
        json={"response": "Yes"}, headers=tokens["donor1"],
    )
    assert response.status_code == 200, response.text
    confirmed = client.post(
        f"/api/admin/donations/requests/{request_id}/donors/{donor_id}/confirm",
        headers=tokens["admin"],
    )
    assert confirmed.status_code == 200, confirmed.text


def test_request_match_response_confirmation_rewards_and_certificate(system):
    client, sessions, tokens = system
    created = client.post("/api/blood-requests", json=blood_request(), headers=tokens["admin"])
    assert created.status_code == 201, created.text
    request_id = created.json()["id"]
    matched = client.post("/api/match/find", json={"blood_request_id": request_id}, headers=tokens["admin"])
    assert matched.status_code == 200 and matched.json()["total_matches"] == 2
    response = client.post(f"/api/donor-dashboard/requests/{request_id}/response", json={"response": "Yes"}, headers=tokens["donor1"])
    assert response.status_code == 200 and response.json()["points_awarded"] == 0
    confirmation = f"/api/admin/donations/requests/{request_id}/donors/1/confirm"
    assert client.post(confirmation, headers=tokens["donor1"]).status_code == 403
    confirmed = client.post(confirmation, headers=tokens["admin"])
    assert confirmed.status_code == 200 and confirmed.json()["points_awarded"] == 100
    repeated = client.post(confirmation, headers=tokens["admin"])
    assert repeated.status_code == 200 and repeated.json()["already_confirmed"]
    with sessions() as db:
        assert len(db.scalars(select(DonationHistory)).all()) == 1
    certificates = client.get("/api/donor-dashboard/certificates", headers=tokens["donor1"])
    assert certificates.status_code == 200 and len(certificates.json()) == 1
    download_url = certificates.json()[0]["download_url"]
    pdf = client.get(download_url, headers=tokens["donor1"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.get(download_url, headers=tokens["donor2"]).status_code in {403, 404}
    assert client.get("/api/donation-history", headers=tokens["admin"]).status_code == 200
    for kind in ("excel", "pdf"):
        assert client.get(f"/api/donation-history/export/{kind}", headers=tokens["admin"]).status_code == 200


def test_donation_history_filters_are_available(system):
    client, _, tokens = system
    response = client.get("/api/donation-history", headers=tokens["admin"])
    assert response.status_code == 200, response.text
    assert response.json()["filters"] == {
        "blood_groups": [],
        "districts": [],
    }


def test_excel_import_export(system):
    client, _, tokens = system
    workbook = Workbook()
    workbook.active.append(["Full Name", "Phone", "Blood Group", "Email"])
    workbook.active.append(["Import Test", "9999900040", "O+", "import@example.org"])
    data = BytesIO(); workbook.save(data)
    response = client.post("/api/donors/import", files={"file": ("donors.xlsx", data.getvalue())}, headers=tokens["admin"])
    assert response.status_code == 200 and response.json()["imported"] == 1, response.text
    output = client.get("/api/donors/export", headers=tokens["admin"])
    assert output.status_code == 200
    assert load_workbook(BytesIO(output.content)).active.max_row == 4


def test_password_reset_single_use_and_old_sessions_revoked(system):
    client, sessions, tokens = system
    token = create_password_reset_token("donor1", 0)
    data = {"token": token, "new_password": "ChangedPassword123"}
    assert client.post("/api/auth/password-reset/confirm", json=data).status_code == 200
    assert client.post("/api/auth/password-reset/confirm", json=data).status_code == 400
    assert client.get("/api/auth/me", headers=tokens["donor1"]).status_code == 401
    assert client.post("/api/auth/login", data={"username": "donor1", "password": "ChangedPassword123"}).status_code == 200


def test_alembic_on_encrypted_database(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + str(tmp_path / "migration-encrypted.db"))
    monkeypatch.setenv("DATABASE_ENCRYPTION_KEY", secrets.token_hex(32))
    get_settings.cache_clear()
    try:
        command.upgrade(Config("alembic.ini"), "head")
        engine = create_database_engine(get_settings())
        with engine.connect() as connection:
            assert connection.execute(text("SELECT version_num FROM alembic_version")).scalar() == "d7e8f9a0b1c2"
            assert connection.execute(text("SELECT COUNT(*) FROM security_rate_limits")).scalar() == 0
        engine.dispose()
    finally:
        get_settings.cache_clear()
